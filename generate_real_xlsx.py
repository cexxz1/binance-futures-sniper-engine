import urllib.request, json, time, datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

TOP_15_120D = [
    'ESPUSDT', 'ONTUSDT', 'MORPHOUSDT', 'MOVRUSDT', 'SEIUSDT',
    'VELVETUSDT', 'ZROUSDT', 'CRVUSDT', 'POLUSDT', 'ZECUSDT',
    'AAVEUSDT', 'CAKEUSDT', 'WLDUSDT', 'ARBUSDT', 'PENDLEUSDT'
]

def fetch_json(url):
    req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
    return json.loads(urllib.request.urlopen(req, timeout=10).read().decode())

def calc_ema(closes, period):
    k = 2 / (period + 1)
    e = closes[0]
    res = [e]
    for c in closes[1:]:
        e = c * k + e * (1 - k)
        res.append(e)
    return res

def calc_rsi(closes, period=14):
    gains, losses = [], []
    for i in range(1, len(closes)):
        d = closes[i] - closes[i-1]
        gains.append(max(d, 0))
        losses.append(max(-d, 0))
    res = [50.0] * period
    ag = sum(gains[:period]) / period
    al = sum(losses[:period]) / period
    for i in range(period, len(gains)):
        ag = (ag * (period - 1) + gains[i]) / period
        al = (al * (period - 1) + losses[i]) / period
        rs = ag / al if al > 0 else 100
        res.append(100 - 100 / (1 + rs))
    res.append(res[-1])
    return res

all_candles = {}
for s in TOP_15_120D:
    k2 = fetch_json(f'https://fapi.binance.com/fapi/v1/klines?symbol={s}&interval=1h&limit=1440')
    first_ts = k2[0][0]
    k1 = fetch_json(f'https://fapi.binance.com/fapi/v1/klines?symbol={s}&interval=1h&limit=1440&endTime={first_ts - 1}')
    all_candles[s] = k1 + k2

min_len = min(len(all_candles[s]) for s in TOP_15_120D)
timestamps = [all_candles['MORPHOUSDT'][i][0] for i in range(min_len)]

indicators = {}
for s in TOP_15_120D:
    closes = [float(k[4]) for k in all_candles[s][:min_len]]
    highs  = [float(k[2]) for k in all_candles[s][:min_len]]
    lows   = [float(k[3]) for k in all_candles[s][:min_len]]
    vols   = [float(k[5]) for k in all_candles[s][:min_len]]
    e9 = calc_ema(closes, 9)
    e21 = calc_ema(closes, 21)
    e99 = calc_ema(closes, 99)
    r14 = calc_rsi(closes, 14)
    indicators[s] = {'closes': closes, 'highs': highs, 'lows': lows, 'vols': vols, 'e9': e9, 'e21': e21, 'e99': e99, 'r14': r14}

balance = 120.0
leverage = 10
pos = None
trades = []

for i in range(100, min_len):
    ts = timestamps[i]
    dt = datetime.datetime.fromtimestamp(ts/1000, datetime.timezone.utc)
    
    if pos is not None:
        s = pos['sym']
        cur_price = indicators[s]['closes'][i]
        cur_high = indicators[s]['highs'][i]
        cur_low = indicators[s]['lows'][i]
        e9 = indicators[s]['e9'][i]
        e21 = indicators[s]['e21'][i]
        
        pos['peak'] = max(pos['peak'], cur_high)
        gain_peak = (pos['peak'] - pos['entry']) / pos['entry']
        
        if pos['pyr_level'] == 0 and gain_peak >= 0.03:
            pos['margin'] += pos['base_margin'] * 0.25
            pos['pyr_level'] = 1
        elif pos['pyr_level'] == 1 and gain_peak >= 0.06:
            pos['margin'] += pos['base_margin'] * 0.25
            pos['pyr_level'] = 2
            
        sl = pos['entry'] * 0.980
        if gain_peak >= 0.25: sl = max(sl, pos['peak'] * 0.95)
        elif gain_peak >= 0.15: sl = max(sl, pos['entry'] * 1.10)
        elif gain_peak >= 0.08: sl = max(sl, pos['entry'] * 1.03)
        elif gain_peak >= 0.03: sl = max(sl, pos['entry'] * 1.002)
        
        exit_trade = False
        exit_p = cur_price
        reason = ''
        if cur_low <= sl:
            exit_trade = True
            exit_p = sl
            reason = 'Dinamik Stop Loss (SL)'
        elif e9 < e21:
            exit_trade = True
            exit_p = cur_price
            reason = '1H Kapanis EMA9 < EMA21'
            
        if exit_trade:
            rg = (exit_p - pos['entry']) / pos['entry']
            pnl = pos['margin'] * leverage * (rg - 0.002)
            balance += pnl
            
            trades.append({
                'id': len(trades) + 1,
                'symbol': s,
                'entry_time': str(pos['entry_dt'])[:19] + " UTC",
                'exit_time': str(dt)[:19] + " UTC",
                'entry_price': pos['entry'],
                'exit_price': exit_p,
                'duration_hours': int((ts - pos['entry_ts'])/3600000),
                'margin_usd': pos['margin'],
                'notional_usd': pos['margin'] * leverage,
                'pnl_usd': pnl,
                'pnl_tl': pnl * 48.0,
                'pnl_pct': rg * leverage * 100,
                'balance_usd': balance,
                'balance_tl': balance * 48.0,
                'reason': reason
            })
            pos = None
            
    if pos is None:
        if dt.hour in [0, 2, 14, 19, 20]: continue
        candidates = []
        for s in TOP_15_120D:
            e9_c = indicators[s]['e9'][i]
            e21_c = indicators[s]['e21'][i]
            e9_p = indicators[s]['e9'][i-1]
            e21_p = indicators[s]['e21'][i-1]
            e99_c = indicators[s]['e99'][i]
            rsi = indicators[s]['r14'][i] if i < len(indicators[s]['r14']) else 50
            vols = indicators[s]['vols']
            v_sma = sum(vols[i-20:i]) / 20.0 if i >= 20 else 1.0
            v_ratio = vols[i] / v_sma if v_sma > 0 else 1.0
            
            if (e9_p <= e21_p) and (e9_c > e21_c) and (e9_c > e99_c) and (rsi >= 54) and (v_ratio >= 1.4):
                candidates.append((s, rsi, v_ratio))
        if candidates:
            candidates.sort(key=lambda x: (x[1], x[2]), reverse=True)
            chosen = candidates[0][0]
            pos = {
                'sym': chosen, 'entry': indicators[chosen]['closes'][i],
                'entry_ts': ts, 'entry_dt': dt, 'peak': indicators[chosen]['closes'][i],
                'base_margin': balance, 'margin': balance, 'pyr_level': 0
            }

# Gerçek Excel (.xlsx) oluştur
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "120 Gunluk Borsa Islemleri"

# Başlık stili
header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
border_thin = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

headers = [
    'İşlem No', 'Parite', 'Giriş Tarihi (UTC)', 'Çıkış Tarihi (UTC)',
    'Giriş Fiyatı ($)', 'Çıkış Fiyatı ($)', 'Süre (Saat)',
    'Kullanılan Marjin ($)', '10x Pozisyon Boyutu ($)', 'İşlem PnL ($)',
    'İşlem PnL (TL)', 'Net Getiri (%)', 'İşlem Sonu Kasa ($)', 'İşlem Sonu Kasa (TL)', 'Kapanış Sebebi'
]

ws.append(headers)
for col_num in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col_num)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center")

for r_idx, t in enumerate(trades, start=2):
    pnl_font_color = "008000" if t['pnl_usd'] >= 0 else "C00000"
    pnl_font = Font(name="Arial", size=10, bold=True, color=pnl_font_color)
    row_data = [
        t['id'],
        t['symbol'],
        t['entry_time'],
        t['exit_time'],
        t['entry_price'],
        t['exit_price'],
        t['duration_hours'],
        t['margin_usd'],
        t['notional_usd'],
        t['pnl_usd'],
        t['pnl_tl'],
        f"%{t['pnl_pct']:.2f}",
        t['balance_usd'],
        t['balance_tl'],
        t['reason']
    ]
    ws.append(row_data)
    
    # Formatlama
    for c_idx in range(1, len(row_data) + 1):
        cell = ws.cell(row=r_idx, column=c_idx)
        cell.border = border_thin
        if c_idx in [1, 7]:
            cell.alignment = Alignment(horizontal="center")
        elif c_idx in [3, 4]:
            cell.alignment = Alignment(horizontal="center")
            cell.font = Font(name="Consolas", size=9)
        elif c_idx in [5, 6]:
            cell.number_format = '#,##0.00000' if t['entry_price'] < 1 else '#,##0.00'
        elif c_idx in [8, 9, 10, 13]:
            cell.number_format = '$#,##0.00'
        elif c_idx in [11, 14]:
            cell.number_format = '#,##0 "TL"'
            
        if c_idx in [10, 11, 12]:
            cell.font = pnl_font

# Otomatik sütun genişliği
for col in ws.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

xlsx_path = 'C:/Users/depco/OneDrive/Desktop/120_GUNLUK_BINANCE_ISLEMLER.xlsx'
wb.save(xlsx_path)
print(f"Gercek XLSX dosyasi olusturuldu: {xlsx_path}")
