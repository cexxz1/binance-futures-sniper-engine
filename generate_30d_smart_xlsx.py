import urllib.request, json, datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

TOP_15_120D = [
    'ESPUSDT', 'ONTUSDT', 'MORPHOUSDT', 'MOVRUSDT', 'SEIUSDT',
    'VELVETUSDT', 'ZROUSDT', 'CRVUSDT', 'POLUSDT', 'ZECUSDT',
    'AAVEUSDT', 'CAKEUSDT', 'WLDUSDT', 'ARBUSDT', 'PENDLEUSDT'
]

def fetch_json(url):
    req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
    return json.loads(urllib.request.urlopen(req, timeout=10).read().decode())

def calc_ema(closes, period):
    k = 2 / (period + 1)
    e = closes[0]
    res = [e]
    for c in closes[1:]:
        e = c * k + e * (1 - k)
        res.append(e)
    return res

def calc_rsi(closes, period=14):
    gains, losses = [], []
    for i in range(1, len(closes)):
        d = closes[i] - closes[i-1]
        gains.append(max(d, 0))
        losses.append(max(-d, 0))
    res = [50.0] * period
    ag = sum(gains[:period]) / period
    al = sum(losses[:period]) / period
    for i in range(period, len(gains)):
        ag = (ag * (period - 1) + gains[i]) / period
        al = (al * (period - 1) + losses[i]) / period
        rs = ag / al if al > 0 else 100
        res.append(100 - 100 / (1 + rs))
    res.append(res[-1])
    return res

def calc_atr(highs, lows, closes, period=14):
    trs = [highs[0] - lows[0]]
    for i in range(1, len(closes)):
        tr = max(highs[i] - lows[i], abs(highs[i] - closes[i-1]), abs(lows[i] - closes[i-1]))
        trs.append(tr)
    atr = [sum(trs[:period]) / period]
    for i in range(period, len(trs)):
        atr.append((atr[-1] * (period - 1) + trs[i]) / period)
    padding = [atr[0]] * (len(closes) - len(atr))
    return padding + atr

all_candles = {}
for s in TOP_15_120D:
    klines = fetch_json(f'https://fapi.binance.com/fapi/v1/klines?symbol={s}&interval=1h&limit=720')
    all_candles[s] = klines

min_len = min(len(all_candles[s]) for s in TOP_15_120D)
timestamps = [all_candles['MORPHOUSDT'][i][0] for i in range(min_len)]

indicators = {}
for s in TOP_15_120D:
    closes = [float(k[4]) for k in all_candles[s][:min_len]]
    highs  = [float(k[2]) for k in all_candles[s][:min_len]]
    lows   = [float(k[3]) for k in all_candles[s][:min_len]]
    vols   = [float(k[5]) for k in all_candles[s][:min_len]]
    indicators[s] = {
        'closes': closes, 'highs': highs, 'lows': lows, 'vols': vols,
        'e9': calc_ema(closes, 9),
        'e21': calc_ema(closes, 21),
        'e99': calc_ema(closes, 99),
        'r14': calc_rsi(closes, 14),
        'atr': calc_atr(highs, lows, closes, 14)
    }

balance = 120.0
pos = None
trades = []

for i in range(50, min_len):
    ts = timestamps[i]
    dt = datetime.datetime.fromtimestamp(ts/1000, datetime.timezone.utc)
    
    if pos is not None:
        s = pos['sym']
        cur_price = indicators[s]['closes'][i]
        cur_high = indicators[s]['highs'][i]
        cur_low = indicators[s]['lows'][i]
        e9 = indicators[s]['e9'][i]
        e21 = indicators[s]['e21'][i]
        lev = pos['leverage']
        
        pos['peak'] = max(pos['peak'], cur_high)
        gain_peak = (pos['peak'] - pos['entry']) / pos['entry']
        
        # Piramit
        if pos['pyr_level'] == 0 and gain_peak >= 0.03:
            pos['margin'] += pos['base_margin'] * 0.25
            pos['pyr_level'] = 1
        elif pos['pyr_level'] == 1 and gain_peak >= 0.06:
            pos['margin'] += pos['base_margin'] * 0.25
            pos['pyr_level'] = 2
            
        # Dinamik Koruma SL
        sl = pos['entry'] * 0.980
        if gain_peak >= 0.25: sl = max(sl, pos['peak'] * 0.95)
        elif gain_peak >= 0.15: sl = max(sl, pos['entry'] * 1.10)
        elif gain_peak >= 0.08: sl = max(sl, pos['entry'] * 1.03)
        elif gain_peak >= 0.03: sl = max(sl, pos['entry'] * 1.002)
        
        exit_trade = False
        exit_p = cur_price
        reason = ''
        if cur_low <= sl:
            exit_trade = True
            exit_p = sl
            reason = 'Trailing SL / Kar Kilitleme'
        elif e9 < e21:
            exit_trade = True
            exit_p = cur_price
            reason = '1H Kapanis EMA9 < EMA21'
            
        if exit_trade:
            rg = (exit_p - pos['entry']) / pos['entry']
            fee = 0.0005 * 2
            pnl = pos['margin'] * lev * (rg - fee)
            balance += pnl
            
            trades.append({
                'sym': s, 'entry_dt': pos['entry_dt'].strftime('%d.%m.%Y %H:%M'), 
                'exit_dt': dt.strftime('%d.%m.%Y %H:%M'),
                'entry_p': pos['entry'], 'exit_p': exit_p, 'lev': lev,
                'used_margin': pos['margin'], 'pos_size': pos['margin'] * lev,
                'pnl': pnl, 'balance': balance, 'reason': reason, 'gain_pct': rg * 100
            })
            pos = None
            
    if pos is None and balance > 5:
        if dt.hour in [0, 2, 14, 19, 20]: continue
        candidates = []
        for s in TOP_15_120D:
            e9_c = indicators[s]['e9'][i]
            e21_c = indicators[s]['e21'][i]
            e9_p = indicators[s]['e9'][i-1]
            e21_p = indicators[s]['e21'][i-1]
            e99_c = indicators[s]['e99'][i]
            rsi = indicators[s]['r14'][i] if i < len(indicators[s]['r14']) else 50
            vols = indicators[s]['vols']
            v_sma = sum(vols[i-20:i]) / 20.0 if i >= 20 else 1.0
            v_ratio = vols[i] / v_sma if v_sma > 0 else 1.0
            atr = indicators[s]['atr'][i]
            atr_pct = (atr / indicators[s]['closes'][i]) * 100
            
            if (e9_p <= e21_p) and (e9_c > e21_c) and (e9_c > e99_c) and (rsi >= 54) and (v_ratio >= 1.4):
                candidates.append((s, rsi, v_ratio, atr_pct))
        if candidates:
            candidates.sort(key=lambda x: (x[1], x[2]), reverse=True)
            chosen, rsi_val, vol_val, atr_val = candidates[0]
            
            # AKILLI KALDIRAÇ ÖLÇÜM MOTORU:
            if atr_val <= 1.5 and vol_val >= 2.0 and rsi_val >= 60:
                calc_lev = 25
            elif atr_val <= 2.2 and vol_val >= 1.6 and rsi_val >= 56:
                calc_lev = 18
            elif atr_val <= 3.2:
                calc_lev = 14
            else:
                calc_lev = 10
                
            pos = {
                'sym': chosen, 'entry': indicators[chosen]['closes'][i],
                'entry_ts': ts, 'entry_dt': dt, 'peak': indicators[chosen]['closes'][i],
                'base_margin': balance, 'margin': balance, 'pyr_level': 0,
                'leverage': calc_lev
            }

# EXCEL DOSYASI ÜRETİMİ
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "30G Akilli Kaldirac Raporu"

headers = [
    "İşlem No", "Parite", "Giriş Tarihi", "Çıkış Tarihi", "Kaldıraç",
    "Giriş Fiyatı ($)", "Çıkış Fiyatı ($)", "Kullanılan Marjin ($)",
    "Pozisyon Boyutu ($)", "İşlem PnL ($)", "İşlem PnL (TL)",
    "Getiri (%)", "İşlem Sonu Kasa ($)", "İşlem Sonu Kasa (TL)", "Çıkış Nedeni"
]

header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
border_thin = Border(
    left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9')
)

ws.append(headers)
for col in range(1, len(headers)+1):
    cell = ws.cell(row=1, column=col)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

for idx, t in enumerate(trades, 1):
    row = [
        idx, t['sym'], t['entry_dt'], t['exit_dt'], f"{t['lev']}x",
        t['entry_p'], t['exit_p'], t['used_margin'], t['pos_size'],
        t['pnl'], t['pnl'] * 48.0, t['gain_pct'],
        t['balance'], t['balance'] * 48.0, t['reason']
    ]
    ws.append(row)
    r_idx = idx + 1
    
    pnl_cell = ws.cell(row=r_idx, column=10)
    if t['pnl'] >= 0:
        pnl_cell.fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
        pnl_cell.font = Font(name="Arial", color="375623", bold=True)
    else:
        pnl_cell.fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
        pnl_cell.font = Font(name="Arial", color="C65911", bold=True)
        
    for c in range(1, len(headers)+1):
        ws.cell(row=r_idx, column=c).border = border_thin

for col in ws.columns:
    max_l = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = max(max_l + 3, 12)

out_path = "C:/Users/depco/OneDrive/Desktop/30_GUNLUK_AKILLI_KALDIRAC_RAPORU.xlsx"
wb.save(out_path)
print(f"XLSX Basarili: {out_path}")
